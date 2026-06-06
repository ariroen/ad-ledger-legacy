#!/usr/bin/env python3
import json
import re
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill


WORKBOOK = Path("/Users/UserOne/Downloads/Закупка рекламы Май Обнова 1 (1).xlsx")
TRACKING_JSON = Path("/tmp/kontrakt_tracking_links_500.json")
REPORT = Path("/Users/UserOne/Documents/New project 2/outputs/view_update_report_20260522.json")

CODE_RE = re.compile(r"/go/([A-Za-z0-9_-]+)")


def norm(value):
    return str(value).strip().lower() if value is not None else ""


def extract_code(value):
    if value is None:
        return None
    text = str(value).strip()
    match = CODE_RE.search(text)
    if match:
        return match.group(1)
    if re.fullmatch(r"[A-Za-z0-9_-]{6,16}", text):
        return text
    return None


def find_view_columns(ws):
    pairs = []
    for row in range(1, ws.max_row + 1):
        ref_cols = [col for col in range(1, ws.max_column + 1) if norm(ws.cell(row, col).value) == "рефка"]
        if not ref_cols:
            continue
        for ref_col in ref_cols:
            next_ref = min([c for c in ref_cols if c > ref_col], default=ws.max_column + 1)
            view_col = None
            for col in range(ref_col + 1, min(next_ref, ws.max_column + 1)):
                if norm(ws.cell(row, col).value) in {"переходы", "просмотры"}:
                    view_col = col
                    break
            if view_col:
                pairs.append((row, ref_col, view_col))
    return pairs


def main():
    payload = json.loads(TRACKING_JSON.read_text(encoding="utf-8"))
    links = payload["data"]
    by_code = {str(item["code"]): item for item in links if item.get("code")}

    wb = load_workbook(WORKBOOK)
    updated = []
    missing_codes = []

    changed_fill = PatternFill("solid", fgColor="E2F0D9")

    for ws in wb.worksheets:
        if ws.title == "Переходы":
            for row in range(1, ws.max_row + 1):
                code = extract_code(ws.cell(row, 2).value)
                if not code:
                    continue
                item = by_code.get(code)
                if not item:
                    missing_codes.append({"sheet": ws.title, "row": row, "code": code})
                    continue
                cell = ws.cell(row, 5)
                old = cell.value
                new = int(item.get("visits") or 0)
                if old != new:
                    cell.value = new
                    cell.fill = changed_fill
                    updated.append({"sheet": ws.title, "cell": cell.coordinate, "code": code, "old": old, "new": new})

        for header_row, ref_col, view_col in find_view_columns(ws):
            for row in range(header_row + 1, ws.max_row + 1):
                # Stop if the same block starts again lower on the sheet.
                if norm(ws.cell(row, ref_col).value) == "рефка":
                    break
                code = extract_code(ws.cell(row, ref_col).value)
                if not code:
                    continue
                item = by_code.get(code)
                if not item:
                    missing_codes.append({"sheet": ws.title, "row": row, "code": code})
                    continue
                cell = ws.cell(row, view_col)
                old = cell.value
                new = int(item.get("visits") or 0)
                if old != new:
                    cell.value = new
                    cell.fill = changed_fill
                    updated.append({"sheet": ws.title, "cell": cell.coordinate, "code": code, "old": old, "new": new})

    wb.save(WORKBOOK)

    report = {
        "updated_at": datetime.now().isoformat(timespec="seconds"),
        "workbook": str(WORKBOOK),
        "api_generated_at": payload.get("generated_at"),
        "api_links": len(links),
        "updated_cells": len(updated),
        "missing_codes": len(missing_codes),
        "updates": updated,
        "missing": missing_codes,
    }
    REPORT.parent.mkdir(parents=True, exist_ok=True)
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({k: report[k] for k in ["api_links", "updated_cells", "missing_codes"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
