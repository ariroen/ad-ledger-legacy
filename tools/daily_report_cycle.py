#!/usr/bin/env python3
import json
import subprocess
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path

from docx import Document

PROJECT_ROOT = Path("/Users/UserOne/Documents/New project 2")
SCREENS_ROOT = PROJECT_ROOT / "outputs/daily_ad_screens"

DEFAULT_XLSX = Path("/Users/UserOne/Downloads/Таблицы/Закупка рекламы Май Обнова 1 (1).xlsx")


@dataclass(frozen=True)
class RunResult:
    ok: bool
    cmd: list[str]
    returncode: int
    stdout: str
    stderr: str


def _run(cmd: list[str]) -> RunResult:
    p = subprocess.run(cmd, check=False, capture_output=True, text=True)
    return RunResult(
        ok=(p.returncode == 0),
        cmd=cmd,
        returncode=p.returncode,
        stdout=(p.stdout or "").strip(),
        stderr=(p.stderr or "").strip(),
    )


def main() -> int:
    today = datetime.now().date()
    yesterday = today - timedelta(days=1)

    xlsx = DEFAULT_XLSX
    if not xlsx.exists():
        raise SystemExit(f"XLSX not found: {xlsx}")

    # 1) Generate/refresh screenshots+manifest for today AND yesterday.
    ad_run = PROJECT_ROOT / "tools/ad_run.py"
    r1 = _run(
        [
            "python3",
            str(ad_run),
            "--xlsx",
            str(xlsx),
            "--out-root",
            str(SCREENS_ROOT),
            "--date",
            today.isoformat(),
            "--date",
            yesterday.isoformat(),
        ]
    )
    if not r1.ok:
        # Still continue to report generation if today folder/manifest already exists.
        pass

    # 2) Build today's report:
    # - include missing/blocked for TODAY, so you always get screenshots;
    # - allow late items, but ONLY when they are found (so "не вышло" doesn't bloat tomorrow).
    report = PROJECT_ROOT / "tools/generate_daily_ad_report.py"
    r2 = _run(
        [
            "python3",
            str(report),
            "--date",
            today.isoformat(),
            "--include-missing",
            "--include-blocked",
            "--late-found-only",
        ]
    )

    verify = {"ok": False, "reason": "report_build_failed"}
    if r2.ok:
        try:
            meta = json.loads(r2.stdout)
            out_path = Path(str(meta.get("out_path") or ""))
            if not out_path.exists():
                raise RuntimeError(f"docx_not_found:{out_path}")
            doc = Document(str(out_path))
            texts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            required = [
                "Количество вышедших постов в ТГ и МАХ",
                "Количество входящих звонков по рекламе",
                "Количество исходящих звонков по ЮВО",
                "Внесено в ГИЦ",
                "Количество людей, готовых прибыть на пункт",
                "Количество людей, прозвоненных по линии СПО",
                "Количество входящих звонков по рекламе УК",
                "Общее количество проработанных звонков за день",
                "Приложение: скриншоты размещений",
            ]
            missing = [t for t in required if not any(t in p for p in texts)]
            image_rels = [r for r in doc.part._rels.values() if "image" in r.reltype]
            verify = {
                "ok": len(missing) == 0,
                "out_path": str(out_path),
                "missing_sections": missing,
                "tables": len(doc.tables),
                "image_rels": len(image_rels),
            }
        except Exception as e:
            verify = {"ok": False, "reason": f"verify_error:{e}"}

    out = {
        "today": today.isoformat(),
        "yesterday": yesterday.isoformat(),
        "xlsx": str(xlsx),
        "screens_root": str(SCREENS_ROOT),
        "ad_run": r1.__dict__,
        "report_build": r2.__dict__,
        "verify": verify,
    }
    print(json.dumps(out, ensure_ascii=False, indent=2))
    return 0 if r2.ok else 2


if __name__ == "__main__":
    raise SystemExit(main())
